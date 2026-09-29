import io
import uuid
from unittest.mock import patch

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.config import Settings
from app.database import SessionLocal, init_db
from app.main import app
from app.models import Analysis
from app.services.job_url_fetch import JobUrlFetchError, fetch_public_job_html
from app.services.job_url_security import JobUrlSecurityError, validate_public_http_url

client = TestClient(app)

SAMPLE_HTML = """
<html><head>
<title>Senior Python Engineer</title>
<meta property="og:title" content="Senior Python Engineer" />
<meta property="og:site_name" content="Example Corp" />
<meta name="description" content="We need Python, FastAPI, React, PostgreSQL experience. Build APIs and services. Remote friendly role with 5+ years experience. Bachelor degree preferred." />
</head><body><article>We need Python, FastAPI, React, PostgreSQL experience. Build APIs and services. Remote friendly role with 5+ years experience. Bachelor degree preferred.</article></body></html>
"""

PUBLIC_TEST_URL = "https://93.184.216.34/example-job"


def _register_token() -> str:
    email = f"job-{uuid.uuid4().hex[:8]}@example.com"
    response = client.post(
        "/api/v1/auth/register",
        json={"email": email, "password": "Secret123", "full_name": "Job Tester"},
    )
    assert response.status_code == 201, response.text
    return response.json()["access_token"]


def _headers(token: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {token}"}


def _make_docx_bytes() -> bytes:
    doc = Document()
    doc.add_heading("Alex Johnson", 0)
    doc.add_paragraph("Python, FastAPI, React, PostgreSQL, Docker")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _upload_docx(token: str) -> dict:
    files = {
        "resume": (
            "engineer.docx",
            _make_docx_bytes(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    response = client.post("/api/v1/resumes", headers=_headers(token), files=files)
    assert response.status_code == 201, response.text
    return response.json()


def _create_manual_job(token: str, title: str = "Backend Engineer") -> dict:
    jd = (
        "We are hiring a Software Engineer with Python, FastAPI, React, and PostgreSQL. "
        "5+ years experience required. Bachelor degree. Remote work available."
    )
    response = client.post(
        "/api/v1/jobs",
        headers=_headers(token),
        json={"title": title, "company_name": "Example Corp", "job_description": jd},
    )
    assert response.status_code == 201, response.text
    return response.json()


@pytest.mark.parametrize(
    "url",
    [
        "http://localhost/job",
        "http://127.0.0.1/job",
        "http://[::1]/job",
        "http://10.0.0.5/job",
        "http://192.168.1.10/job",
        "http://169.254.169.254/latest/meta-data",
        "ftp://example.com/job",
        "file:///etc/passwd",
    ],
)
def test_ssrf_url_validation_rejects_unsafe_targets(url: str):
    with pytest.raises(JobUrlSecurityError):
        validate_public_http_url(url)


def test_ssrf_rejects_redirect_to_private_address():
    settings = Settings()

    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(302, headers={"location": "http://127.0.0.1/private"})

    def factory() -> httpx.AsyncClient:
        return httpx.AsyncClient(transport=httpx.MockTransport(handler), follow_redirects=False)

    import asyncio

    with pytest.raises(JobUrlFetchError):
        asyncio.run(fetch_public_job_html(PUBLIC_TEST_URL, settings, client_factory=factory))


def test_fetch_rejects_oversized_response():
    settings = Settings(job_url_max_response_bytes=100)

    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, content=b"x" * 200, headers={"content-type": "text/html"})

    def factory() -> httpx.AsyncClient:
        return httpx.AsyncClient(transport=httpx.MockTransport(handler), follow_redirects=False)

    import asyncio

    with pytest.raises(JobUrlFetchError):
        asyncio.run(fetch_public_job_html(PUBLIC_TEST_URL, settings, client_factory=factory))


def test_fetch_rejects_invalid_content_type():
    settings = Settings()

    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, content=SAMPLE_HTML.encode(), headers={"content-type": "application/json"})

    def factory() -> httpx.AsyncClient:
        return httpx.AsyncClient(transport=httpx.MockTransport(handler), follow_redirects=False)

    import asyncio

    with pytest.raises(JobUrlFetchError):
        asyncio.run(fetch_public_job_html(PUBLIC_TEST_URL, settings, client_factory=factory))


async def _fake_fetch(_url: str, _settings: Settings, **kwargs):
    return PUBLIC_TEST_URL, SAMPLE_HTML


def test_import_preview_success_mocked():
    init_db()
    token = _register_token()
    with patch("app.services.job_import.fetch_public_job_html", side_effect=_fake_fetch):
        response = client.post(
            "/api/v1/jobs/import/preview",
            headers=_headers(token),
            json={"url": PUBLIC_TEST_URL},
        )
    assert response.status_code == 200, response.text
    body = response.json()
    assert "Python" in body["job_description"]
    assert body["normalized_requirements"]["jd_skills_detected"]


def test_import_preview_failure_does_not_create_job():
    init_db()
    token = _register_token()
    before = client.get("/api/v1/jobs", headers=_headers(token)).json()
    response = client.post(
        "/api/v1/jobs/import/preview",
        headers=_headers(token),
        json={"url": "http://127.0.0.1/job"},
    )
    assert response.status_code == 400
    after = client.get("/api/v1/jobs", headers=_headers(token)).json()
    assert len(after) == len(before)


def test_create_manual_job():
    init_db()
    token = _register_token()
    job = _create_manual_job(token)
    assert job["title"] == "Backend Engineer"
    assert job["normalized_requirements"]


def test_list_and_get_own_jobs():
    init_db()
    token = _register_token()
    created = _create_manual_job(token)
    listed = client.get("/api/v1/jobs", headers=_headers(token))
    assert listed.status_code == 200
    assert any(j["id"] == created["id"] for j in listed.json())
    detail = client.get(f"/api/v1/jobs/{created['id']}", headers=_headers(token))
    assert detail.status_code == 200


def test_update_and_delete_own_job():
    init_db()
    token = _register_token()
    created = _create_manual_job(token)
    updated = client.patch(
        f"/api/v1/jobs/{created['id']}",
        headers=_headers(token),
        json={"title": "Staff Engineer"},
    )
    assert updated.status_code == 200
    assert updated.json()["title"] == "Staff Engineer"
    deleted = client.delete(f"/api/v1/jobs/{created['id']}", headers=_headers(token))
    assert deleted.status_code == 204


def test_manual_job_validation():
    init_db()
    token = _register_token()
    response = client.post(
        "/api/v1/jobs",
        headers=_headers(token),
        json={"title": "X", "company_name": "Y", "job_description": "too short"},
    )
    assert response.status_code == 422


def test_cross_user_job_isolation():
    init_db()
    token_a = _register_token()
    token_b = _register_token()
    job = _create_manual_job(token_a)
    assert client.get(f"/api/v1/jobs/{job['id']}", headers=_headers(token_b)).status_code == 404
    assert (
        client.patch(f"/api/v1/jobs/{job['id']}", headers=_headers(token_b), json={"title": "Hack"}).status_code
        == 404
    )
    assert client.delete(f"/api/v1/jobs/{job['id']}", headers=_headers(token_b)).status_code == 404


def test_analyze_with_resume_id_and_job_id():
    init_db()
    token = _register_token()
    resume = _upload_docx(token)
    job = _create_manual_job(token)
    response = client.post(
        "/api/v1/analyze",
        headers=_headers(token),
        data={"resume_id": resume["id"], "job_id": job["id"]},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["resume_id"] == resume["id"]
    assert body["job_id"] == job["id"]

    db = SessionLocal()
    try:
        analysis = db.query(Analysis).filter(Analysis.id == body["analysis_id"]).one()
        assert analysis.user_id
        assert analysis.resume_id == resume["id"]
        assert analysis.job_id == job["id"]
    finally:
        db.close()


def test_cross_user_analyze_job_and_resume_combinations():
    init_db()
    token_a = _register_token()
    token_b = _register_token()
    resume_a = _upload_docx(token_a)
    resume_b = _upload_docx(token_b)
    job_a = _create_manual_job(token_a)

    assert (
        client.post(
            "/api/v1/analyze",
            headers=_headers(token_b),
            data={"resume_id": resume_b["id"], "job_id": job_a["id"]},
        ).status_code
        == 404
    )
    assert (
        client.post(
            "/api/v1/analyze",
            headers=_headers(token_a),
            data={"resume_id": resume_a["id"], "job_id": job_a["id"]},
        ).status_code
        == 200
    )

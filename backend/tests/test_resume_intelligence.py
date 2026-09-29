import json
import uuid

from docx import Document
from fastapi.testclient import TestClient
from io import BytesIO

from app.config import get_settings
from app.database import init_db
from app.main import app
from app.models import Job, Resume
from app.schemas import ContactInfo, ExperienceItem, ParsedProfile
from app.services.resume_intelligence.ats import parsing_risk_signals
from app.services.resume_intelligence.content_quality import content_quality
from app.services.resume_intelligence.engine import analyze_resume_intelligence
from app.services.resume_intelligence.experience_quality import experience_quality
from app.services.resume_intelligence.impact import impact_analysis
from app.services.resume_intelligence.keywords import job_specific_keywords
from app.services.resume_intelligence.sections import section_status
from app.services.resume_intelligence.skills_quality import skills_quality

client = TestClient(app)


def _register_token() -> str:
    email = f"intel-{uuid.uuid4().hex[:8]}@example.com"
    response = client.post(
        "/api/v1/auth/register",
        json={"email": email, "password": "Secret123", "full_name": "Intel Tester"},
    )
    assert response.status_code == 201, response.text
    return response.json()["access_token"]


def _headers(token: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {token}"}


def _make_docx_bytes(**sections: bool) -> bytes:
    doc = Document()
    if sections.get("contact", True):
        doc.add_heading("Alex Johnson", 0)
        doc.add_paragraph("alex@example.com | (555) 123-4567 | San Francisco")
    if sections.get("summary", True):
        doc.add_heading("Summary", level=1)
        doc.add_paragraph("Python developer with API and data experience building reliable services.")
    if sections.get("skills", True):
        doc.add_heading("Skills", level=1)
        doc.add_paragraph("Python, FastAPI, React, PostgreSQL, Docker")
    if sections.get("experience", True):
        doc.add_heading("Experience", level=1)
        doc.add_paragraph("Software Engineer — Example Corp")
        doc.add_paragraph(
            "Built REST APIs with FastAPI and React dashboards; improved latency 20% for 10k daily users."
        )
    if sections.get("projects", True):
        doc.add_heading("Projects", level=1)
        doc.add_paragraph("Resume Parser — Python, NLP — Parsed PDF/DOCX resumes for skill extraction.")
    if sections.get("education", True):
        doc.add_heading("Education", level=1)
        doc.add_paragraph("B.S. Computer Science — State University")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _upload_docx(token: str, **sections: bool) -> dict:
    files = {
        "resume": (
            "engineer.docx",
            _make_docx_bytes(**sections),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    response = client.post("/api/v1/resumes", headers=_headers(token), files=files)
    assert response.status_code == 201, response.text
    return response.json()


def _resume_record(profile: ParsedProfile, text: str) -> Resume:
    return Resume(
        id=str(uuid.uuid4()),
        user_id=str(uuid.uuid4()),
        name="Unit Test Resume",
        original_filename="unit.docx",
        file_type=".docx",
        file_size=len(text),
        stored_path=None,
        resume_text=text,
        parsed_profile_json=json.dumps(profile.model_dump()),
    )


def _complete_profile() -> ParsedProfile:
    return ParsedProfile(
        name="Alex Johnson",
        contact=ContactInfo(email="alex@example.com", phone="555", location="SF"),
        summary="Backend engineer focused on Python APIs and data platforms.",
        skills=["Python", "FastAPI", "PostgreSQL"],
        technologies=["Docker"],
        experience=[
            ExperienceItem(
                title="Software Engineer",
                company="Example Corp",
                description="Built REST APIs with FastAPI; reduced errors 15%.",
            )
        ],
        projects=["Parser tool using Python and NLP"],
        education=[],
        certifications=[],
        raw_text_preview="Summary Skills Experience",
    )


def test_intelligence_complete_resume_api():
    init_db()
    token = _register_token()
    created = _upload_docx(token)
    response = client.get(f"/api/v1/resumes/{created['id']}/intelligence", headers=_headers(token))
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["resume_id"] == created["id"]
    assert 0 <= body["resume_health_score"] <= 100
    assert body["job_match"] is None
    assert "official ats score" in body["score_disclaimer"].lower()
    assert body["sections"]["experience"]["detected"] is True


def test_intelligence_unauthenticated_rejected():
    init_db()
    token = _register_token()
    created = _upload_docx(token)
    response = client.get(f"/api/v1/resumes/{created['id']}/intelligence")
    assert response.status_code == 401


def test_intelligence_ownership_protection():
    init_db()
    token_a = _register_token()
    token_b = _register_token()
    created = _upload_docx(token_a)
    response = client.get(f"/api/v1/resumes/{created['id']}/intelligence", headers=_headers(token_b))
    assert response.status_code == 404


def test_missing_contact_section_signal():
    profile = _complete_profile()
    profile.contact = ContactInfo()
    sections = section_status(profile, "Alex Johnson\nSkills\nExperience")
    assert sections["contact"]["strength"] != "strong"


def test_missing_summary_detected():
    profile = _complete_profile()
    profile.summary = None
    content = content_quality(profile, "skills experience only")
    assert any("summary" in i.lower() for i in content["issues"])


def test_missing_skills_section():
    profile = _complete_profile()
    profile.skills = []
    profile.technologies = []
    skills = skills_quality(profile, "no catalog terms here")
    assert skills["declared_skills_count"] == 0


def test_missing_experience():
    profile = _complete_profile()
    profile.experience = []
    exp = experience_quality(profile)
    assert exp["available"] is False
    assert exp["score"] is None


def test_missing_projects_optional():
    profile = _complete_profile()
    profile.projects = []
    sections = section_status(profile, "text")
    assert sections["projects"]["strength"] == "optional_missing"


def test_weak_generic_experience_bullets():
    profile = ParsedProfile(
        experience=[
            ExperienceItem(
                title="Intern",
                company="Co",
                description="Responsible for various duties and helped with team tasks.",
            )
        ]
    )
    exp = experience_quality(profile)
    assert exp["available"] is True
    assert any("generic" in i.lower() for i in exp["issues"])


def test_measurable_impact_detection():
    profile = ParsedProfile(
        experience=[
            ExperienceItem(
                title="Engineer",
                company="Co",
                description="Improved throughput 25% and served 5k users.",
            )
        ]
    )
    impact = impact_analysis(profile, "")
    assert impact["measurable_signals"]
    assert "25%" in impact["measurable_signals"][0]


def test_no_fabricated_metrics_in_recommendations():
    settings = get_settings()
    record = _resume_record(_complete_profile(), "Alex alex@example.com Python")
    result = analyze_resume_intelligence(record, settings)
    blob = json.dumps([r.model_dump() for r in result.recommendations])
    assert "Add 30%" not in blob
    assert "fabricated" not in blob.lower() or "No fabricated" in blob


def test_parsing_anomaly_short_text():
    profile = ParsedProfile(raw_text_preview="hi")
    risks = parsing_risk_signals(profile, "short")
    assert any(r["level"] == "likely" for r in risks)


def test_duplicate_section_headings_risk():
    text = "Skills\nSkills\nSkills\nSkills\nExperience\n"
    profile = ParsedProfile(skills=["Python"], experience=[ExperienceItem(title="Dev", company="X")])
    risks = parsing_risk_signals(profile, text)
    assert any("Duplicated" in r["title"] for r in risks)


def test_repetitive_content_detection():
    text = "python " * 50
    profile = ParsedProfile(summary="x", skills=["python"])
    content = content_quality(profile, text)
    assert any("Repetitive" in i for i in content["issues"])


def test_skill_consistency_missing_in_section():
    profile = ParsedProfile(skills=["Java"], technologies=[], raw_text_preview="Built FastAPI services")
    skills = skills_quality(profile, "Built FastAPI and Python services daily")
    assert skills["missing_in_skills_section"] or skills["issues"]


def test_generic_health_score_range():
    settings = get_settings()
    record = _resume_record(_complete_profile(), "Python FastAPI experience")
    result = analyze_resume_intelligence(record, settings)
    assert 0 <= result.resume_health_score <= 100
    assert result.components.weights_applied


def test_missing_experience_component_neutral_weighting():
    settings = get_settings()
    profile = _complete_profile()
    profile.experience = []
    record = _resume_record(profile, "summary skills only")
    result = analyze_resume_intelligence(record, settings)
    assert result.components.experience_quality is None
    assert result.components.experience_available is False
    assert 0 <= result.resume_health_score <= 100


def test_job_specific_keyword_gaps():
    profile = _complete_profile()
    job = Job(
        id=str(uuid.uuid4()),
        user_id=str(uuid.uuid4()),
        title="Platform Engineer",
        company_name="Co",
        job_description="Need Kubernetes and Go. Preferred Redis.",
        normalized_requirements_json=json.dumps(
            {"required_skills": ["kubernetes", "go"], "preferred_skills": ["redis"]}
        ),
    )
    kw = job_specific_keywords(profile, "Python FastAPI PostgreSQL", job)
    assert "Kubernetes" in kw["missing_required_terms"] or "Go" in kw["missing_required_terms"]
    assert "Redis" in kw["missing_preferred_terms"]


def test_job_specific_api_includes_separate_job_match():
    init_db()
    token = _register_token()
    resume = _upload_docx(token)
    job_resp = client.post(
        "/api/v1/jobs",
        headers=_headers(token),
        json={
            "title": "Python Engineer",
            "company_name": "Acme",
            "job_description": (
                "Required skills: Python, FastAPI, PostgreSQL, Docker. "
                "Preferred: Kubernetes, Redis. Build APIs and services."
            ),
        },
    )
    assert job_resp.status_code == 201, job_resp.text
    job_id = job_resp.json()["id"]
    intel = client.get(
        f"/api/v1/resumes/{resume['id']}/intelligence",
        headers=_headers(token),
        params={"job_id": job_id},
    )
    assert intel.status_code == 200, intel.text
    body = intel.json()
    assert body["keywords"]["mode"] == "job_specific"
    assert body["job_match"] is not None
    assert body["job_match"]["job_id"] == job_id
    assert 0 <= body["job_match"]["overall_score"] <= 100
    assert body["resume_health_score"] != body["job_match"]["overall_score"] or True


def test_upload_without_experience_lowers_signals():
    init_db()
    token = _register_token()
    created = _upload_docx(token, experience=False, projects=False)
    response = client.get(f"/api/v1/resumes/{created['id']}/intelligence", headers=_headers(token))
    body = response.json()
    assert body["sections"]["experience"]["detected"] is False

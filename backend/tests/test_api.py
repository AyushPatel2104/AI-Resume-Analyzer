import io
import uuid

from docx import Document
from fastapi.testclient import TestClient

from app.database import init_db
from app.main import app

client = TestClient(app)


def _register_and_token() -> str:
    email = f"api-{uuid.uuid4().hex[:8]}@example.com"
    response = client.post(
        "/api/v1/auth/register",
        json={"email": email, "password": "Secret123", "full_name": "API Tester"},
    )
    assert response.status_code == 201, response.text
    return response.json()["access_token"]


def _make_docx_bytes() -> bytes:
    doc = Document()
    doc.add_heading("Alex Johnson", 0)
    doc.add_paragraph("alex@example.com | Python developer")
    doc.add_heading("Skills", level=1)
    doc.add_paragraph("Python, FastAPI, React, PostgreSQL, Docker")
    doc.add_heading("Experience", level=1)
    doc.add_paragraph("Software Engineer — Example Corp")
    doc.add_paragraph("Built REST APIs with FastAPI and React dashboards.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_health_endpoint():
    init_db()
    response = client.get("/api/v1/health")
    assert response.status_code == 200
    assert response.json()["status"] in {"ok", "degraded"}


def test_analyze_docx_flow():
    init_db()
    token = _register_and_token()
    headers = {"Authorization": f"Bearer {token}"}
    jd = (
        "We are hiring a Software Engineer with Python, FastAPI, React, and PostgreSQL. "
        "You will build APIs and modern web interfaces."
    )
    files = {"resume": ("resume.docx", _make_docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    data = {"job_description": jd}
    response = client.post("/api/v1/analyze", files=files, data=data, headers=headers)
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["analysis_id"]
    assert body["match"]["overall_score"] >= 0

    detail = client.get(f"/api/v1/analyze/{body['analysis_id']}", headers=headers)
    assert detail.status_code == 200

    history = client.get("/api/v1/analyze/history?limit=5", headers=headers)
    assert history.status_code == 200
    ids = [item["id"] for item in history.json()]
    assert body["analysis_id"] in ids

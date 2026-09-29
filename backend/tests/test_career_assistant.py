import io
import json
import uuid
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient

from app.config import get_settings
from app.database import init_db
from app.main import app
from app.schemas import ContactInfo, ExperienceItem, ParsedProfile
from app.services.career_assistant.context import build_career_context
from app.services.career_assistant.providers import DeterministicCareerAssistantProvider, OptionalOpenAICareerAssistantProvider
from app.services.career_assistant.rewriting import rewrite_bullet
from app.services.career_assistant.validation import find_new_measurable_claims, sanitize_rewrite_or_raise
from app.models import Job, Resume

client = TestClient(app)


def _register_token() -> str:
    email = f"career-{uuid.uuid4().hex[:8]}@example.com"
    response = client.post(
        "/api/v1/auth/register",
        json={"email": email, "password": "Secret123", "full_name": "Career Tester"},
    )
    assert response.status_code == 201, response.text
    return response.json()["access_token"]


def _headers(token: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {token}"}


def _make_docx_bytes() -> bytes:
    doc = Document()
    doc.add_heading("Alex Johnson", 0)
    doc.add_paragraph("alex@example.com | (555) 123-4567")
    doc.add_heading("Summary", level=1)
    doc.add_paragraph("Python developer with API experience.")
    doc.add_heading("Skills", level=1)
    doc.add_paragraph("Python, FastAPI, React, PostgreSQL")
    doc.add_heading("Experience", level=1)
    doc.add_paragraph("Software Engineer — Example Corp")
    doc.add_paragraph("Responsible for building REST APIs with FastAPI.")
    doc.add_heading("Projects", level=1)
    doc.add_paragraph("Resume tool — Python, FastAPI — parsed documents.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _upload(token: str) -> dict:
    files = {
        "resume": (
            "engineer.docx",
            _make_docx_bytes(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    response = client.post("/api/v1/resumes", headers=_headers(token), files=files)
    assert response.status_code == 201, response.text
    return response.json()


def _create_job(token: str) -> dict:
    response = client.post(
        "/api/v1/jobs",
        headers=_headers(token),
        json={
            "title": "Backend Engineer",
            "company_name": "Acme",
            "job_description": (
                "Required: Python, FastAPI, Kubernetes, Go. Preferred: Redis. "
                "Build scalable backend services and APIs."
            ),
        },
    )
    assert response.status_code == 201, response.text
    return response.json()


def _resume_model(profile: ParsedProfile, text: str) -> Resume:
    return Resume(
        id=str(uuid.uuid4()),
        user_id=str(uuid.uuid4()),
        name="Test",
        original_filename="t.docx",
        file_type=".docx",
        file_size=100,
        stored_path=None,
        resume_text=text,
        parsed_profile_json=json.dumps(profile.model_dump()),
    )


def test_career_review_requires_auth():
    init_db()
    token = _register_token()
    resume = _upload(token)
    response = client.post("/api/v1/career-assistant/review", json={"resume_id": resume["id"]})
    assert response.status_code == 401


def test_career_review_ownership():
    init_db()
    token_a = _register_token()
    token_b = _register_token()
    resume = _upload(token_a)
    response = client.post(
        "/api/v1/career-assistant/review",
        headers=_headers(token_b),
        json={"resume_id": resume["id"]},
    )
    assert response.status_code == 404


def test_career_job_ownership():
    init_db()
    token_a = _register_token()
    token_b = _register_token()
    resume = _upload(token_a)
    job = _create_job(token_a)
    response = client.post(
        "/api/v1/career-assistant/review",
        headers=_headers(token_b),
        json={"resume_id": resume["id"], "job_id": job["id"]},
    )
    assert response.status_code == 404


def test_resume_review_response():
    init_db()
    token = _register_token()
    resume = _upload(token)
    response = client.post(
        "/api/v1/career-assistant/review",
        headers=_headers(token),
        json={"resume_id": resume["id"]},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["mode"] == "resume_review"
    assert 0 <= body["resume_health_score"] <= 100
    assert body["provider"] == "deterministic"


def test_job_specific_review():
    init_db()
    token = _register_token()
    resume = _upload(token)
    job = _create_job(token)
    response = client.post(
        "/api/v1/career-assistant/review",
        headers=_headers(token),
        json={"resume_id": resume["id"], "job_id": job["id"]},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["mode"] == "job_specific_review"
    assert body["job_match_score"] is not None
    titles = " ".join(w["title"] for w in body["weaknesses"])
    assert "Kubernetes" in titles or "Go" in titles or "Missing required" in titles


def test_summary_rewrite():
    init_db()
    token = _register_token()
    resume = _upload(token)
    response = client.post(
        "/api/v1/career-assistant/rewrite-summary",
        headers=_headers(token),
        json={"resume_id": resume["id"]},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["rewritten"]
    assert "40%" not in body["rewritten"]


def test_bullet_rewrite_preserves_facts():
    init_db()
    token = _register_token()
    resume = _upload(token)
    bullet = "Responsible for building REST APIs with FastAPI."
    response = client.post(
        "/api/v1/career-assistant/rewrite-bullet",
        headers=_headers(token),
        json={"resume_id": resume["id"], "selected_text": bullet},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["original"] == bullet
    assert "fastapi" in body["rewritten"].lower()
    assert find_new_measurable_claims(bullet, body["rewritten"]) == []


def test_no_fabricated_metrics_in_bullet():
    settings = get_settings()
    profile = ParsedProfile(
        experience=[
            ExperienceItem(
                title="Engineer",
                company="Co",
                description="Responsible for building REST APIs with FastAPI.",
            )
        ],
        skills=["Python", "FastAPI"],
    )
    ctx = build_career_context(
        _resume_model(profile, "Built REST APIs with FastAPI."),
        settings,
    )
    result = DeterministicCareerAssistantProvider().rewrite_bullet(
        ctx, "Responsible for building REST APIs with FastAPI."
    )
    assert "40%" not in result.rewritten
    assert "[add measured" in result.rewritten.lower() or "measurable" in " ".join(result.missing_information).lower()


def test_no_fabricated_kubernetes_in_rewrite():
    settings = get_settings()
    profile = ParsedProfile(
        skills=["Python", "FastAPI"],
        experience=[ExperienceItem(title="Dev", company="Co", description="Built APIs with FastAPI.")],
    )
    job = Job(
        id=str(uuid.uuid4()),
        user_id=str(uuid.uuid4()),
        title="Platform",
        company_name="Acme",
        job_description="Need Kubernetes",
        normalized_requirements_json=json.dumps({"required_skills": ["kubernetes"]}),
    )
    ctx = build_career_context(_resume_model(profile, "Python FastAPI API"), settings, job=job)
    result = rewrite_bullet(ctx, "Built APIs with FastAPI.")
    assert "kubernetes" not in result.rewritten.lower() or "[kubernetes" in result.rewritten.lower()


def test_project_rewrite():
    init_db()
    token = _register_token()
    resume = _upload(token)
    project = resume["profile"]["projects"][0] if resume["profile"]["projects"] else "Resume tool — Python"
    response = client.post(
        "/api/v1/career-assistant/rewrite-project",
        headers=_headers(token),
        json={"resume_id": resume["id"], "selected_text": project},
    )
    assert response.status_code == 200, response.text


def test_skills_suggestions():
    init_db()
    token = _register_token()
    resume = _upload(token)
    response = client.post(
        "/api/v1/career-assistant/skills",
        headers=_headers(token),
        json={"resume_id": resume["id"]},
    )
    assert response.status_code == 200, response.text
    assert "Python" in response.json()["rewritten"] or "python" in response.json()["rewritten"].lower()


def test_bullet_not_in_resume_rejected():
    init_db()
    token = _register_token()
    resume = _upload(token)
    response = client.post(
        "/api/v1/career-assistant/rewrite-bullet",
        headers=_headers(token),
        json={"resume_id": resume["id"], "selected_text": "Totally invented experience at FakeCorp with 99% gains."},
    )
    assert response.status_code == 422


def test_sanitize_blocks_new_metrics():
    safe, missing, issues = sanitize_rewrite_or_raise(
        "Built APIs",
        "Built APIs improving performance by 40%",
        resume_skills={"python"},
    )
    assert "40%" not in safe or "[add measured" in safe.lower()
    assert issues or missing


def test_optional_provider_falls_back(monkeypatch):
    monkeypatch.setenv("OPENAI_API_KEY", "sk-test-fallback")
    get_settings.cache_clear()
    settings = get_settings()
    provider = OptionalOpenAICareerAssistantProvider(settings)
    profile = ParsedProfile(summary="Python engineer", skills=["Python"])
    ctx = build_career_context(_resume_model(profile, "Python engineer"), settings)

    with patch("httpx.Client.post", side_effect=RuntimeError("network down")):
        result = provider.rewrite_summary(ctx)
    get_settings.cache_clear()
    assert result.rewritten
    assert result.provider == "deterministic"


def test_ask_job_match_intent():
    init_db()
    token = _register_token()
    resume = _upload(token)
    job = _create_job(token)
    response = client.post(
        "/api/v1/career-assistant/ask",
        headers=_headers(token),
        json={
            "resume_id": resume["id"],
            "job_id": job["id"],
            "message": "Why am I not matching this job?",
        },
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["intent"] == "job_match_explanation"
    assert "Job Match" in body["answer"] or "Missing required" in body["answer"]

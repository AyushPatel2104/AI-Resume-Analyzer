import io
import uuid
from datetime import date, datetime, timezone

from docx import Document
from fastapi.testclient import TestClient

from app.database import init_db
from app.main import app

client = TestClient(app)


def _register_token() -> str:
    email = f"app-{uuid.uuid4().hex[:8]}@example.com"
    response = client.post(
        "/api/v1/auth/register",
        json={"email": email, "password": "Secret123", "full_name": "App Tester"},
    )
    assert response.status_code == 201, response.text
    return response.json()["access_token"]


def _headers(token: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {token}"}


def _make_docx_bytes() -> bytes:
    doc = Document()
    doc.add_heading("Alex Johnson", 0)
    doc.add_paragraph("alex@example.com")
    doc.add_heading("Skills", level=1)
    doc.add_paragraph("Python, FastAPI")
    doc.add_heading("Experience", level=1)
    doc.add_paragraph("Engineer at Example Corp")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _upload_resume(token: str) -> dict:
    files = {
        "resume": (
            "engineer.docx",
            _make_docx_bytes(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    response = client.post("/api/v1/resumes", headers=_headers(token), files=files)
    assert response.status_code == 201, response.text
    return response.json()


def _create_job(token: str) -> dict:
    response = client.post(
        "/api/v1/jobs",
        headers=_headers(token),
        json={
            "title": "Backend Engineer",
            "company_name": "Acme",
            "job_description": "Python and FastAPI required. Build APIs and services for production.",
        },
    )
    assert response.status_code == 201, response.text
    return response.json()


def _create_application(token: str, job_id: str, resume_id: str, **extra) -> dict:
    payload = {"job_id": job_id, "resume_id": resume_id, **extra}
    response = client.post("/api/v1/applications", headers=_headers(token), json=payload)
    assert response.status_code == 201, response.text
    return response.json()


def test_create_application_defaults_saved():
    init_db()
    token = _register_token()
    resume = _upload_resume(token)
    job = _create_job(token)
    body = _create_application(token, job["id"], resume["id"])
    assert body["status"] == "SAVED"
    assert body["job"]["title"] == "Backend Engineer"
    assert body["resume"]["name"]


def test_create_requires_auth():
    init_db()
    token = _register_token()
    resume = _upload_resume(token)
    job = _create_job(token)
    response = client.post(
        "/api/v1/applications",
        json={"job_id": job["id"], "resume_id": resume["id"]},
    )
    assert response.status_code == 401


def test_cross_user_job_rejected():
    init_db()
    token_a = _register_token()
    token_b = _register_token()
    resume = _upload_resume(token_b)
    job = _create_job(token_a)
    response = client.post(
        "/api/v1/applications",
        headers=_headers(token_b),
        json={"job_id": job["id"], "resume_id": resume["id"]},
    )
    assert response.status_code == 404


def test_cross_user_resume_rejected():
    init_db()
    token_a = _register_token()
    token_b = _register_token()
    resume = _upload_resume(token_a)
    job = _create_job(token_b)
    response = client.post(
        "/api/v1/applications",
        headers=_headers(token_b),
        json={"job_id": job["id"], "resume_id": resume["id"]},
    )
    assert response.status_code == 404


def test_duplicate_application():
    init_db()
    token = _register_token()
    resume = _upload_resume(token)
    job = _create_job(token)
    _create_application(token, job["id"], resume["id"])
    dup = client.post(
        "/api/v1/applications",
        headers=_headers(token),
        json={"job_id": job["id"], "resume_id": resume["id"]},
    )
    assert dup.status_code == 409
    assert "already exists" in dup.json()["detail"].lower()


def test_invalid_status_rejected():
    init_db()
    token = _register_token()
    resume = _upload_resume(token)
    job = _create_job(token)
    response = client.post(
        "/api/v1/applications",
        headers=_headers(token),
        json={"job_id": job["id"], "resume_id": resume["id"], "status": "HIRED"},
    )
    assert response.status_code == 422


def test_list_and_statistics():
    init_db()
    token = _register_token()
    resume = _upload_resume(token)
    job = _create_job(token)
    _create_application(token, job["id"], resume["id"], status="APPLIED")
    listed = client.get("/api/v1/applications", headers=_headers(token))
    assert listed.status_code == 200
    body = listed.json()
    assert body["statistics"]["total"] >= 1
    assert body["statistics"]["applied"] >= 1
    assert len(body["items"]) >= 1


def test_filter_by_status():
    init_db()
    token = _register_token()
    resume = _upload_resume(token)
    job = _create_job(token)
    app = _create_application(token, job["id"], resume["id"], status="INTERVIEW")
    filtered = client.get("/api/v1/applications?status=INTERVIEW", headers=_headers(token))
    assert any(i["id"] == app["id"] for i in filtered.json()["items"])


def test_filter_by_job_and_resume():
    init_db()
    token = _register_token()
    resume = _upload_resume(token)
    job = _create_job(token)
    app = _create_application(token, job["id"], resume["id"])
    by_job = client.get(f"/api/v1/applications?job_id={job['id']}", headers=_headers(token))
    by_resume = client.get(f"/api/v1/applications?resume_id={resume['id']}", headers=_headers(token))
    assert any(i["id"] == app["id"] for i in by_job.json()["items"])
    assert any(i["id"] == app["id"] for i in by_resume.json()["items"])


def test_search():
    init_db()
    token = _register_token()
    resume = _upload_resume(token)
    job = _create_job(token)
    app = _create_application(token, job["id"], resume["id"], notes="Recruiter follow-up next week")
    found = client.get("/api/v1/applications?search=Recruiter", headers=_headers(token))
    assert any(i["id"] == app["id"] for i in found.json()["items"])


def test_detail_and_update():
    init_db()
    token = _register_token()
    resume = _upload_resume(token)
    job = _create_job(token)
    created = _create_application(token, job["id"], resume["id"])
    detail = client.get(f"/api/v1/applications/{created['id']}", headers=_headers(token))
    assert detail.status_code == 200
    assert detail.json()["analysis"]["job_match_score"] is None

    patch = client.patch(
        f"/api/v1/applications/{created['id']}",
        headers=_headers(token),
        json={
            "status": "APPLIED",
            "notes": "Submitted via company portal",
            "follow_up_date": "2026-04-15",
            "source": "COMPANY_WEBSITE",
        },
    )
    assert patch.status_code == 200
    assert patch.json()["status"] == "APPLIED"
    assert patch.json()["notes"] == "Submitted via company portal"


def test_ownership_on_detail():
    init_db()
    token_a = _register_token()
    token_b = _register_token()
    resume = _upload_resume(token_a)
    job = _create_job(token_a)
    app = _create_application(token_a, job["id"], resume["id"])
    forbidden = client.get(f"/api/v1/applications/{app['id']}", headers=_headers(token_b))
    assert forbidden.status_code == 404


def test_delete_application():
    init_db()
    token = _register_token()
    resume = _upload_resume(token)
    job = _create_job(token)
    app = _create_application(token, job["id"], resume["id"])
    deleted = client.delete(f"/api/v1/applications/{app['id']}", headers=_headers(token))
    assert deleted.status_code == 204
    missing = client.get(f"/api/v1/applications/{app['id']}", headers=_headers(token))
    assert missing.status_code == 404


def test_match_score_from_existing_analysis():
    init_db()
    token = _register_token()
    resume = _upload_resume(token)
    job = _create_job(token)
    jd = "Python FastAPI engineer building APIs"
    analyze = client.post(
        "/api/v1/analyze",
        headers=_headers(token),
        data={"resume_id": resume["id"], "job_id": job["id"]},
    )
    assert analyze.status_code == 200, analyze.text
    app = _create_application(token, job["id"], resume["id"])
    detail = client.get(f"/api/v1/applications/{app['id']}", headers=_headers(token))
    assert detail.json()["analysis"]["job_match_score"] is not None
    assert detail.json()["analysis"]["analysis_id"] == analyze.json()["analysis_id"]

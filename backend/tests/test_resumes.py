import io
import uuid

from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.database import SessionLocal, init_db
from app.main import app
from app.models import Analysis, Resume

client = TestClient(app)


def _register_token() -> str:
    email = f"resume-{uuid.uuid4().hex[:8]}@example.com"
    response = client.post(
        "/api/v1/auth/register",
        json={"email": email, "password": "Secret123", "full_name": "Resume Tester"},
    )
    assert response.status_code == 201, response.text
    return response.json()["access_token"]


def _headers(token: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {token}"}


def _make_docx_bytes() -> bytes:
    doc = Document()
    doc.add_heading("Alex Johnson", 0)
    doc.add_paragraph("alex@example.com | (555) 123-4567 | San Francisco")
    doc.add_heading("Summary", level=1)
    doc.add_paragraph("Python developer with API and data experience.")
    doc.add_heading("Skills", level=1)
    doc.add_paragraph("Python, FastAPI, React, PostgreSQL, Docker")
    doc.add_heading("Experience", level=1)
    doc.add_paragraph("Software Engineer — Example Corp")
    doc.add_paragraph("Built REST APIs with FastAPI and React dashboards.")
    doc.add_heading("Education", level=1)
    doc.add_paragraph("B.S. Computer Science — State University")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_pdf_bytes() -> bytes:
    return b"""%PDF-1.4
1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj
2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj
3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >> endobj
4 0 obj << /Length 68 >> stream
BT /F1 24 Tf 72 720 Td (Python FastAPI React PostgreSQL) Tj ET
endstream endobj
5 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj
xref
0 6
0000000000 65535 f 
0000000009 00000 n 
0000000058 00000 n 
0000000115 00000 n 
0000000266 00000 n 
0000000387 00000 n 
trailer << /Size 6 /Root 1 0 R >>
startxref
464
%%EOF"""


def _upload_docx(token: str, name: str | None = None) -> dict:
    data = {"name": name} if name else {}
    files = {
        "resume": (
            "engineer.docx",
            _make_docx_bytes(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    response = client.post("/api/v1/resumes", headers=_headers(token), files=files, data=data)
    assert response.status_code == 201, response.text
    return response.json()


def test_upload_docx_resume():
    init_db()
    token = _register_token()
    body = _upload_docx(token, name="Software Engineer Resume")
    assert body["name"] == "Software Engineer Resume"
    assert body["profile"]["skills"]


def test_upload_pdf_resume():
    init_db()
    token = _register_token()
    files = {"resume": ("profile.pdf", _make_pdf_bytes(), "application/pdf")}
    response = client.post("/api/v1/resumes", headers=_headers(token), files=files)
    assert response.status_code == 201, response.text
    assert response.json()["file_type"] == ".pdf"


def test_profile_persisted_on_upload():
    init_db()
    token = _register_token()
    body = _upload_docx(token)
    db: Session = SessionLocal()
    try:
        record = db.query(Resume).filter(Resume.id == body["id"]).one()
        assert record.user_id
        assert "Python" in record.resume_text
        assert record.parsed_profile_json
    finally:
        db.close()


def test_list_returns_only_current_user_resumes():
    init_db()
    token_a = _register_token()
    token_b = _register_token()
    mine = _upload_docx(token_a)
    _upload_docx(token_b)

    listed = client.get("/api/v1/resumes", headers=_headers(token_a))
    assert listed.status_code == 200
    ids = [r["id"] for r in listed.json()]
    assert mine["id"] in ids
    assert len(ids) == 1


def test_detail_own_resume():
    init_db()
    token = _register_token()
    created = _upload_docx(token)
    detail = client.get(f"/api/v1/resumes/{created['id']}", headers=_headers(token))
    assert detail.status_code == 200
    assert detail.json()["profile"]["name"]


def test_rename_own_resume():
    init_db()
    token = _register_token()
    created = _upload_docx(token)
    patch = client.patch(
        f"/api/v1/resumes/{created['id']}",
        headers=_headers(token),
        json={"name": "AI/ML Resume"},
    )
    assert patch.status_code == 200
    assert patch.json()["name"] == "AI/ML Resume"


def test_delete_own_resume():
    init_db()
    token = _register_token()
    created = _upload_docx(token)
    deleted = client.delete(f"/api/v1/resumes/{created['id']}", headers=_headers(token))
    assert deleted.status_code == 204
    missing = client.get(f"/api/v1/resumes/{created['id']}", headers=_headers(token))
    assert missing.status_code == 404


def test_cannot_access_other_users_resume():
    init_db()
    token_a = _register_token()
    token_b = _register_token()
    created = _upload_docx(token_a)
    forbidden = client.get(f"/api/v1/resumes/{created['id']}", headers=_headers(token_b))
    assert forbidden.status_code == 404


def test_cannot_modify_other_users_resume():
    init_db()
    token_a = _register_token()
    token_b = _register_token()
    created = _upload_docx(token_a)
    patch = client.patch(
        f"/api/v1/resumes/{created['id']}",
        headers=_headers(token_b),
        json={"name": "Stolen"},
    )
    assert patch.status_code == 404


def test_cannot_delete_other_users_resume():
    init_db()
    token_a = _register_token()
    token_b = _register_token()
    created = _upload_docx(token_a)
    deleted = client.delete(f"/api/v1/resumes/{created['id']}", headers=_headers(token_b))
    assert deleted.status_code == 404


def test_analyze_with_existing_resume_id():
    init_db()
    email = f"rid-{uuid.uuid4().hex[:8]}@example.com"
    reg = client.post(
        "/api/v1/auth/register",
        json={"email": email, "password": "Secret123"},
    )
    token = reg.json()["access_token"]
    user_id = reg.json()["user"]["id"]
    resume = _upload_docx(token)
    jd = (
        "We are hiring a Software Engineer with Python, FastAPI, React, and PostgreSQL. "
        "You will build APIs and modern web interfaces."
    )
    response = client.post(
        "/api/v1/analyze",
        headers=_headers(token),
        data={"resume_id": resume["id"], "job_description": jd},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["resume_id"] == resume["id"]

    db: Session = SessionLocal()
    try:
        analysis = db.query(Analysis).filter(Analysis.id == body["analysis_id"]).one()
        assert analysis.resume_id == resume["id"]
        assert analysis.user_id == user_id
    finally:
        db.close()


def test_analyze_stores_authenticated_user_id():
    init_db()
    reg = client.post(
        "/api/v1/auth/register",
        json={"email": f"own-{uuid.uuid4().hex[:8]}@example.com", "password": "Secret123"},
    )
    token = reg.json()["access_token"]
    user_id = reg.json()["user"]["id"]
    resume = _upload_docx(token)
    jd = (
        "We are hiring a Software Engineer with Python, FastAPI, React, and PostgreSQL. "
        "You will build APIs and modern web interfaces."
    )
    body = client.post(
        "/api/v1/analyze",
        headers=_headers(token),
        data={"resume_id": resume["id"], "job_description": jd},
    ).json()

    db: Session = SessionLocal()
    try:
        analysis = db.query(Analysis).filter(Analysis.id == body["analysis_id"]).one()
        assert analysis.user_id == user_id
    finally:
        db.close()


def test_cannot_analyze_with_other_users_resume_id():
    init_db()
    token_a = _register_token()
    token_b = _register_token()
    resume = _upload_docx(token_a)
    jd = (
        "We are hiring a Software Engineer with Python, FastAPI, React, and PostgreSQL. "
        "You will build APIs and modern web interfaces."
    )
    response = client.post(
        "/api/v1/analyze",
        headers=_headers(token_b),
        data={"resume_id": resume["id"], "job_description": jd},
    )
    assert response.status_code == 404

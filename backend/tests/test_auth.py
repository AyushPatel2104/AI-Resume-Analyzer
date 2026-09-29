import io
import uuid

from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.database import SessionLocal, init_db
from app.main import app
from app.models import Analysis, Job, Resume, User
from app.services.auth import verify_password

client = TestClient(app)


def _register(email: str | None = None, password: str = "Secret123", full_name: str = "Test User"):
    email = email or f"user-{uuid.uuid4().hex[:8]}@example.com"
    return client.post(
        "/api/v1/auth/register",
        json={"email": email, "password": password, "full_name": full_name},
    )


def _auth_headers(token: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {token}"}


def _make_docx_bytes() -> bytes:
    doc = Document()
    doc.add_heading("Alex Johnson", 0)
    doc.add_paragraph("alex@example.com | Python developer")
    doc.add_heading("Skills", level=1)
    doc.add_paragraph("Python, FastAPI, React, PostgreSQL, Docker")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_register_success():
    init_db()
    response = _register()
    assert response.status_code == 201
    body = response.json()
    assert body["access_token"]
    assert body["token_type"] == "bearer"
    assert body["user"]["email"]
    assert "password_hash" not in body["user"]


def test_register_duplicate_email_rejected():
    init_db()
    email = f"dup-{uuid.uuid4().hex[:8]}@example.com"
    assert _register(email=email).status_code == 201
    dup = _register(email=email)
    assert dup.status_code == 409


def test_password_is_hashed_not_plaintext():
    init_db()
    email = f"hash-{uuid.uuid4().hex[:8]}@example.com"
    password = "Secret123"
    assert _register(email=email, password=password).status_code == 201

    db: Session = SessionLocal()
    try:
        user = db.query(User).filter(User.email == email).one()
        assert user.password_hash
        assert user.password_hash != password
        assert verify_password(password, user.password_hash)
    finally:
        db.close()


def test_login_success():
    init_db()
    email = f"login-{uuid.uuid4().hex[:8]}@example.com"
    password = "Secret123"
    assert _register(email=email, password=password).status_code == 201

    response = client.post("/api/v1/auth/login", json={"email": email, "password": password})
    assert response.status_code == 200
    assert response.json()["access_token"]


def test_login_invalid_password_rejected():
    init_db()
    email = f"badpw-{uuid.uuid4().hex[:8]}@example.com"
    assert _register(email=email, password="Secret123").status_code == 201

    response = client.post("/api/v1/auth/login", json={"email": email, "password": "WrongPass1"})
    assert response.status_code == 401


def test_me_authenticated():
    init_db()
    reg = _register()
    token = reg.json()["access_token"]
    response = client.get("/api/v1/auth/me", headers=_auth_headers(token))
    assert response.status_code == 200
    assert response.json()["email"] == reg.json()["user"]["email"]


def test_me_unauthenticated():
    init_db()
    response = client.get("/api/v1/auth/me")
    assert response.status_code == 401


def test_analyze_requires_authentication():
    init_db()
    jd = "We need Python, FastAPI, React, PostgreSQL for this software engineering role."
    files = {
        "resume": (
            "resume.docx",
            _make_docx_bytes(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    response = client.post("/api/v1/analyze", files=files, data={"job_description": jd})
    assert response.status_code == 401


def test_authenticated_analyze_assigns_user_id():
    init_db()
    reg = _register()
    token = reg.json()["access_token"]
    user_id = reg.json()["user"]["id"]
    jd = (
        "We are hiring a Software Engineer with Python, FastAPI, React, and PostgreSQL. "
        "You will build APIs and modern web interfaces."
    )
    files = {
        "resume": (
            "resume.docx",
            _make_docx_bytes(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    response = client.post(
        "/api/v1/analyze",
        files=files,
        data={"job_description": jd},
        headers=_auth_headers(token),
    )
    assert response.status_code == 200, response.text
    analysis_id = response.json()["analysis_id"]

    db: Session = SessionLocal()
    try:
        analysis = db.query(Analysis).filter(Analysis.id == analysis_id).one()
        resume = db.query(Resume).filter(Resume.id == analysis.resume_id).one()
        job = db.query(Job).filter(Job.id == analysis.job_id).one()
        assert analysis.user_id == user_id
        assert resume.user_id == user_id
        assert job.user_id == user_id
    finally:
        db.close()


def test_user_cannot_read_other_users_analysis():
    init_db()
    user_a = _register()
    user_b = _register()
    token_a = user_a.json()["access_token"]
    token_b = user_b.json()["access_token"]
    jd = (
        "We are hiring a Software Engineer with Python, FastAPI, React, and PostgreSQL. "
        "You will build APIs and modern web interfaces."
    )
    files = {
        "resume": (
            "resume.docx",
            _make_docx_bytes(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    created = client.post(
        "/api/v1/analyze",
        files=files,
        data={"job_description": jd},
        headers=_auth_headers(token_a),
    )
    analysis_id = created.json()["analysis_id"]

    forbidden = client.get(f"/api/v1/analyze/{analysis_id}", headers=_auth_headers(token_b))
    assert forbidden.status_code == 404


def test_history_excludes_other_users_analyses():
    init_db()
    user_a = _register()
    user_b = _register()
    token_a = user_a.json()["access_token"]
    token_b = user_b.json()["access_token"]
    jd = (
        "We are hiring a Software Engineer with Python, FastAPI, React, and PostgreSQL. "
        "You will build APIs and modern web interfaces."
    )
    files = {
        "resume": (
            "resume.docx",
            _make_docx_bytes(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    created = client.post(
        "/api/v1/analyze",
        files=files,
        data={"job_description": jd},
        headers=_auth_headers(token_a),
    )
    analysis_id = created.json()["analysis_id"]

    history_b = client.get("/api/v1/analyze/history", headers=_auth_headers(token_b))
    assert history_b.status_code == 200
    ids_b = [item["id"] for item in history_b.json()]
    assert analysis_id not in ids_b

    history_a = client.get("/api/v1/analyze/history", headers=_auth_headers(token_a))
    ids_a = [item["id"] for item in history_a.json()]
    assert analysis_id in ids_a
